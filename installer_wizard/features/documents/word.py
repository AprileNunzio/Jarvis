import io
import time
from pathlib import Path

from features.documents import charts, themes
from features.documents.spec import number, runs

TONE = {"info": ("2E86DE", "EAF2FB"), "success": ("17A589", "E8F8F5"), "warning": ("F39C12", "FEF5E7"),
        "danger": ("E74C3C", "FDEDEC")}
MONTHS = ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre", "ottobre",
          "novembre", "dicembre"]


def italian_date() -> str:
    now = time.localtime()
    return f"{now.tm_mday} {MONTHS[now.tm_mon - 1]} {now.tm_year}"


def fmt_number(value) -> str:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    if float(value).is_integer():
        return f"{int(value):,}".replace(",", ".")
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


class Writer:
    def __init__(self, doc: dict, links: dict | None = None) -> None:
        from docx import Document
        self.spec = doc
        self.links = links or {}
        self.theme = themes.pick(doc.get("theme"), doc.get("palette"), doc.get("font"))
        self.doc = Document()

    def _oxml(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        return OxmlElement, qn

    def color(self, hex_color: str):
        from docx.shared import RGBColor
        return RGBColor(*themes.rgb(hex_color))

    def shade(self, cell, hex_color: str) -> None:
        OxmlElement, qn = self._oxml()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        cell._tc.get_or_add_tcPr().append(shd)

    def borders(self, cell, color: str = "D5DBDB", left: tuple | None = None, size: int = 4) -> None:
        OxmlElement, qn = self._oxml()
        tc_pr = cell._tc.get_or_add_tcPr()
        box = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            col, sz = (left if edge == "left" and left else (color, size))
            el.set(qn("w:val"), "single" if col else "nil")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), col or "auto")
            box.append(el)
        tc_pr.append(box)

    def setup(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Cm, Pt
        t = self.theme
        for section in self.doc.sections:
            section.page_height, section.page_width = Cm(29.7), Cm(21)
            section.left_margin = section.right_margin = Cm(2.2)
            section.top_margin, section.bottom_margin = Cm(2.2), Cm(2)
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = t["body"], Pt(10.5)
        normal.font.color.rgb = self.color(t["dark"])
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.2
        sizes = {1: 20, 2: 15, 3: 12.5}
        for level, size in sizes.items():
            st = self.doc.styles[f"Heading {level}"]
            st.font.name, st.font.size, st.font.bold = t["heading"], Pt(size), level > 1
            st.font.color.rgb = self.color(t["primary"] if level < 3 else t["accent"])
            st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(18 if level == 1 else 12), Pt(6)
            st.paragraph_format.keep_with_next = True
        if "Didascalia Jarvis" not in [s.name for s in self.doc.styles]:
            cap = self.doc.styles.add_style("Didascalia Jarvis", WD_STYLE_TYPE.PARAGRAPH)
            cap.base_style = normal
            cap.font.size, cap.font.italic = Pt(9), True
            cap.font.color.rgb = self.color(t["muted"])

    def page_field(self, paragraph, instruction: str) -> None:
        from docx.shared import Pt
        OxmlElement, qn = self._oxml()
        run = paragraph.add_run()
        run.font.size = Pt(8.5)
        for kind, value in (("begin", None), ("instr", instruction), ("separate", None), ("end", None)):
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = value
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
            if kind == "separate":
                text = OxmlElement("w:t")
                text.text = "1"
                run._r.append(text)

    def header_footer(self) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        section = self.doc.sections[0]
        section.different_first_page_header_footer = True
        for name in ("Header", "Footer"):
            self.doc.styles[name].font.size = Pt(8.5)
            self.doc.styles[name].font.color.rgb = self.color(self.theme["muted"])
        head = section.header.paragraphs[0]
        head.text = self.spec["title"]
        head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in head.runs:
            r.font.size, r.font.color.rgb = Pt(8.5), self.color(self.theme["muted"])
        foot = section.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        foot.add_run("Pagina ").font.size = Pt(8.5)
        self.page_field(foot, "PAGE")
        foot.add_run(" di ").font.size = Pt(8.5)
        self.page_field(foot, "NUMPAGES")

    def cover(self) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt
        t, s = self.theme, self.spec
        band = self.doc.add_table(rows=1, cols=1)
        cell = band.rows[0].cells[0]
        self.shade(cell, t["primary"])
        self.borders(cell, color="")
        cell.width = Cm(16.6)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(60), Pt(4)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(s["title"])
        r.font.name, r.font.size, r.font.bold = t["heading"], Pt(30), True
        r.font.color.rgb = self.color(themes.readable_on(t["primary"]))
        if s.get("subtitle"):
            p2 = cell.add_paragraph()
            p2.paragraph_format.left_indent = Cm(0.6)
            p2.paragraph_format.space_after = Pt(60)
            r2 = p2.add_run(s["subtitle"])
            r2.font.size = Pt(14)
            r2.font.color.rgb = self.color(themes.readable_on(t["primary"]))
        else:
            p.paragraph_format.space_after = Pt(60)
        meta = self.doc.add_paragraph()
        meta.paragraph_format.space_before = Pt(18)
        meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
        mr = meta.add_run(f"{s.get('author') or 'Jarvis'}  ·  {italian_date()}")
        mr.font.size, mr.font.color.rgb = Pt(10), self.color(t["muted"])
        OxmlElement, qn = self._oxml()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", t["accent"])):
            bottom.set(qn(key), value)
        borders.append(bottom)
        meta._p.get_or_add_pPr().append(borders)

    def hyperlink(self, paragraph, label: str, url: str) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE
        OxmlElement, qn = self._oxml()
        rid = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), self.theme["accent"])
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        props.append(color)
        props.append(underline)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = label
        text.set(qn("xml:space"), "preserve")
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    def rich(self, paragraph, value: str) -> None:
        from docx.shared import Pt
        for chunk, style in runs(value):
            if style.get("link"):
                self.hyperlink(paragraph, chunk, style["link"])
                continue
            r = paragraph.add_run(chunk)
            r.bold = style.get("bold")
            r.italic = style.get("italic")
            if style.get("code"):
                r.font.name, r.font.size = "Consolas", Pt(9.5)
                r.font.color.rgb = self.color(self.theme["accent"])

    def table(self, b: dict) -> None:
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        t = self.theme
        cols, rows = b["columns"], [list(r) for r in b["rows"]]
        numeric = [all(number(r[i]) is not None for r in rows if str(r[i]).strip()) and any(str(r[i]).strip() for r in rows)
                   for i in range(len(cols))]
        if b.get("totals"):
            total = ["Totale"] + [sum(number(r[i]) or 0 for r in rows) if numeric[i] and i else "" for i in range(1, len(cols))]
            rows.append(total)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(cols))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, name in enumerate(cols):
            cell = table.rows[0].cells[i]
            self.shade(cell, t["primary"])
            self.borders(cell, color=t["primary"])
            p = cell.paragraphs[0]
            r = p.add_run(name)
            r.bold, r.font.size = True, Pt(9.5)
            r.font.color.rgb = self.color(themes.readable_on(t["primary"]))
            if numeric[i]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for ri, row in enumerate(rows, start=1):
            last = b.get("totals") and ri == len(rows)
            for ci, value in enumerate(row):
                cell = table.rows[ri].cells[ci]
                self.borders(cell)
                if last:
                    self.shade(cell, t["light"])
                elif ri % 2 == 0:
                    self.shade(cell, "F7F9FA")
                p = cell.paragraphs[0]
                n = number(value) if numeric[ci] else None
                r = p.add_run(fmt_number(n) if n is not None else str(value))
                r.font.size, r.bold = Pt(9.5), bool(last)
                if numeric[ci]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if b.get("caption"):
            self.doc.add_paragraph(b["caption"], style="Didascalia Jarvis")
        else:
            self.doc.add_paragraph()

    def chart(self, b: dict) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm
        try:
            data = charts.png(b, self.theme)
        except Exception as exc:
            self.doc.add_paragraph(f"[Grafico non disponibile: {exc}]", style="Didascalia Jarvis")
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(io.BytesIO(data), width=Cm(15.5))

    def kpi(self, b: dict) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        t = self.theme
        table = self.doc.add_table(rows=1, cols=len(b["items"]))
        for i, item in enumerate(b["items"]):
            cell = table.rows[0].cells[i]
            self.shade(cell, t["light"])
            self.borders(cell, color="FFFFFF", left=(t["accent"], 24))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(item["value"])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(18), self.color(t["primary"])
            lp = cell.add_paragraph()
            lr = lp.add_run(item["label"])
            lr.font.size, lr.font.color.rgb = Pt(9), self.color(t["muted"])
            if item.get("delta"):
                dr = lp.add_run(f"  {item['delta']}")
                negative = item["delta"].strip().startswith(("-", "−"))
                dr.font.size, dr.bold = Pt(9), True
                dr.font.color.rgb = self.color("E74C3C" if negative else "17A589")
        self.doc.add_paragraph()

    def callout(self, b: dict) -> None:
        from docx.shared import Pt
        edge, fill = TONE.get(b.get("tone"), TONE["info"])
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        self.shade(cell, fill)
        self.borders(cell, color=fill, left=(edge, 36))
        p = cell.paragraphs[0]
        if b.get("title"):
            r = p.add_run(b["title"])
            r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), self.color(edge)
            p = cell.add_paragraph()
        self.rich(p, b.get("text", ""))
        self.doc.add_paragraph()

    def quote(self, b: dict) -> None:
        from docx.shared import Cm, Pt
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        self.borders(cell, color="", left=(self.theme["accent"], 24))
        p = cell.paragraphs[0]
        p.paragraph_format.left_indent = Cm(0.3)
        start = len(p.runs)
        self.rich(p, f"«{b['text'].strip('«»')}»")
        for r in p.runs[start:]:
            r.italic, r.font.size, r.font.color.rgb = True, Pt(12), self.color(self.theme["primary"])
        if b.get("author"):
            a = cell.add_paragraph()
            ar = a.add_run(f"— {b['author']}")
            ar.font.size, ar.font.color.rgb = Pt(9), self.color(self.theme["muted"])
        self.doc.add_paragraph()

    def numbered(self, entries: list[str]) -> None:
        from docx.shared import Cm, Pt
        for i, item in enumerate(entries, start=1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.6)
            p.paragraph_format.space_after = Pt(4)
            n = p.add_run(f"{i}.\t")
            n.bold, n.font.color.rgb = True, self.color(self.theme["accent"])
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            self.rich(p, item)

    def blocks(self) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
        for b in self.spec["blocks"]:
            t = b["type"]
            if t == "heading":
                self.doc.add_heading(b["text"], level=b["level"])
            elif t == "paragraph":
                p = self.doc.add_paragraph()
                p.alignment = align.get(b.get("align"), WD_ALIGN_PARAGRAPH.JUSTIFY)
                self.rich(p, b["text"])
            elif t == "bullets":
                for item in b["items"]:
                    self.rich(self.doc.add_paragraph(style="List Bullet"), item)
            elif t == "numbered":
                self.numbered(b["items"])
            elif t == "table":
                self.table(b)
            elif t == "chart":
                self.chart(b)
            elif t == "kpi":
                self.kpi(b)
            elif t == "callout":
                self.callout(b)
            elif t == "quote":
                self.quote(b)
            elif t == "link":
                target = self.links.get(b.get("target", ""))
                if target:
                    p = self.doc.add_paragraph()
                    p.add_run("→ ")
                    self.hyperlink(p, b["text"] or target, target)
            elif t == "pagebreak":
                self.doc.add_page_break()

    def save(self, path: Path) -> Path:
        self.setup()
        self.cover()
        self.doc.add_page_break()
        self.header_footer()
        self.blocks()
        props = self.doc.core_properties
        props.title, props.author, props.subject = self.spec["title"], self.spec.get("author") or "Jarvis", self.spec.get("subtitle", "")
        self.doc.save(str(path))
        return path


def write(doc: dict, path: Path, links: dict | None = None) -> Path:
    return Writer(doc, links).save(path)
