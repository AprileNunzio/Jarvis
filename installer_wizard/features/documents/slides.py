from pathlib import Path

from features.documents import themes
from features.documents.spec import number, runs
from features.documents.word import fmt_number

W, H = 13.333, 7.5


class Deck:
    def __init__(self, doc: dict, links: dict | None = None) -> None:
        from pptx import Presentation
        from pptx.util import Inches
        self.spec, self.links = doc, links or {}
        self.theme = themes.pick(doc.get("theme"), doc.get("palette"), doc.get("font"))
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.count = 0
        self.sections = 0

    def rgb(self, hex_color: str):
        from pptx.dml.color import RGBColor
        return RGBColor(*themes.rgb(hex_color))

    def new(self, background: str | None = None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.count += 1
        if background:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = self.rgb(background)
        return slide

    def box(self, slide, x, y, w, h, color: str, line: str | None = None):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.rgb(color)
        if line:
            shape.line.color.rgb = self.rgb(line)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text(self, slide, x, y, w, h, value: str, size: float, color: str, bold: bool = False, font: str | None = None,
             align: str = "left", italic: bool = False):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        from pptx.util import Inches
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
        self.rich(p, value, size, color, bold, font, italic)
        return tb

    def rich(self, p, value: str, size: float, color: str, bold: bool = False, font: str | None = None, italic: bool = False):
        from pptx.util import Pt
        for chunk, style in runs(value) or [(value, {})]:
            r = p.add_run()
            r.text = chunk
            r.font.size, r.font.name = Pt(size), font or self.theme["body"]
            r.font.bold = bold or style.get("bold")
            r.font.italic = italic or style.get("italic")
            r.font.color.rgb = self.rgb(color)
            target = style.get("link")
            if target:
                r.hyperlink.address = self.links.get(target, target)
                r.font.color.rgb = self.rgb(self.theme["accent"])
                r.font.underline = True

    def bullets(self, slide, x, y, w, h, items: list[str], size: float = 18) -> None:
        from pptx.util import Inches, Pt
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            dot = p.add_run()
            dot.text = "■  "
            dot.font.size, dot.font.color.rgb = Pt(size * 0.6), self.rgb(self.theme["accent"])
            self.rich(p, item, size, self.theme["dark"])

    def chrome(self, slide, title: str) -> None:
        t = self.theme
        self.box(slide, 0, 0, 0.18, H, t["primary"])
        self.text(slide, 0.6, 0.35, W - 1.2, 0.9, title, 30, t["primary"], True, t["heading"])
        self.box(slide, 0.62, 1.22, 1.2, 0.06, t["accent"])
        self.text(slide, 0.6, H - 0.5, 8, 0.3, self.spec["title"], 10, t["muted"])
        self.text(slide, W - 1.6, H - 0.5, 1, 0.3, str(self.count), 10, t["muted"], align="right")

    def table(self, slide, spec: dict, x, y, w, h) -> None:
        from pptx.util import Inches, Pt
        t = self.theme
        cols, rows = spec["columns"], spec["rows"][:14]
        shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(x), Inches(y), Inches(w), Inches(min(h, 0.45 * (len(rows) + 1))))
        table = shape.table
        for i, name in enumerate(cols):
            cell = table.cell(0, i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.rgb(t["primary"])
            cell.text = str(name)
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.bold, run.font.size = True, Pt(13)
            run.font.color.rgb = self.rgb(themes.readable_on(t["primary"]))
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb("F7F9FA" if r % 2 == 0 else "FFFFFF")
                n = number(value) if not isinstance(value, str) or value.replace(",", "").replace(".", "").strip("€% -").isdigit() else None
                cell.text = fmt_number(n) if n is not None else str(value)
                run = cell.text_frame.paragraphs[0].runs[0] if cell.text_frame.paragraphs[0].runs else None
                if run:
                    run.font.size, run.font.color.rgb = Pt(12), self.rgb(t["dark"])

    def chart(self, slide, spec: dict, x, y, w, h) -> None:
        from pptx.chart.data import CategoryChartData, XyChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.util import Inches, Pt
        kind = spec["chart"]
        types = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                 "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
                 "scatter": XL_CHART_TYPE.XY_SCATTER}
        if kind == "scatter":
            data = XyChartData()
            for s in spec["series"]:
                series = data.add_series(s["name"])
                for i, v in enumerate(s["values"]):
                    series.add_data_point(i + 1, v)
        else:
            data = CategoryChartData()
            data.categories = spec["categories"]
            for s in (spec["series"][:1] if kind in ("pie", "doughnut") else spec["series"]):
                data.add_series(s["name"], s["values"])
        frame = slide.shapes.add_chart(types[kind], Inches(x), Inches(y), Inches(w), Inches(h), data)
        chart = frame.chart
        chart.font.size, chart.font.name = Pt(12), self.theme["body"]
        colors = self.theme["chart"]
        if kind in ("pie", "doughnut"):
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.RIGHT
            chart.legend.include_in_layout = False
            for i, point in enumerate(chart.plots[0].series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.rgb(colors[i % len(colors)])
            chart.plots[0].has_data_labels = True
            chart.plots[0].data_labels.number_format = "0%"
            chart.plots[0].data_labels.show_percentage = True
        else:
            for i, series in enumerate(chart.plots[0].series):
                fmt = series.format
                if kind in ("line", "scatter"):
                    fmt.line.color.rgb = self.rgb(colors[i % len(colors)])
                    fmt.line.width = Pt(2.5)
                else:
                    fmt.fill.solid()
                    fmt.fill.fore_color.rgb = self.rgb(colors[i % len(colors)])
            chart.has_legend = len(spec["series"]) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
            if kind in ("bar", "column"):
                chart.plots[0].gap_width = 70
        if spec.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = spec["title"]
            chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)

    def kpis(self, slide, items: list[dict]) -> None:
        t = self.theme
        n = max(1, len(items))
        gap, width = 0.35, (W - 1.2 - 0.35 * (n - 1)) / n
        for i, item in enumerate(items):
            x = 0.6 + i * (width + gap)
            self.box(slide, x, 2.2, width, 3.2, t["light"])
            self.box(slide, x, 2.2, width, 0.09, t["accent"])
            self.text(slide, x + 0.3, 2.7, width - 0.6, 1.3, item["value"], 40, t["primary"], True, t["heading"])
            self.text(slide, x + 0.3, 3.9, width - 0.6, 0.8, item["label"], 16, t["muted"])
            if item.get("delta"):
                negative = item["delta"].strip().startswith(("-", "−"))
                self.text(slide, x + 0.3, 4.6, width - 0.6, 0.5, item["delta"], 16, "E74C3C" if negative else "17A589", True)

    def slide(self, s: dict) -> None:
        t, layout = self.theme, s["layout"]
        if layout in ("title", "closing"):
            slide = self.new(t["primary"])
            ink = themes.readable_on(t["primary"])
            self.box(slide, 0.8, 2.55, 1.4, 0.1, t["accent2"])
            self.text(slide, 0.8, 2.8, W - 1.6, 1.6, s["title"] or self.spec["title"], 44, ink, True, t["heading"])
            self.text(slide, 0.8, 4.4, W - 1.6, 1.2, s["subtitle"] or self.spec.get("subtitle", ""), 20, ink)
            if layout == "title":
                self.text(slide, 0.8, H - 1.0, 8, 0.4, self.spec.get("author") or "Jarvis", 12, ink)
        elif layout == "section":
            slide = self.new("FFFFFF")
            self.sections += 1
            self.box(slide, 0, 0, W * 0.38, H, t["primary"])
            self.text(slide, 0.6, 3.0, W * 0.38 - 1, 1.5, f"{self.sections:02d}", 60, themes.readable_on(t["primary"]), True, t["heading"])
            self.text(slide, W * 0.38 + 0.8, 3.0, W * 0.62 - 1.6, 1.8, s["title"], 38, t["primary"], True, t["heading"])
            if s.get("subtitle"):
                self.text(slide, W * 0.38 + 0.8, 4.4, W * 0.62 - 1.6, 1.0, s["subtitle"], 18, t["muted"])
        elif layout == "quote":
            slide = self.new(t["light"])
            self.text(slide, 1.2, 1.4, W - 2.4, 3.6, f"«{s.get('text') or s.get('title')}»", 34, t["primary"], False, t["heading"], "center", True)
            if s.get("author"):
                self.text(slide, 1.2, 5.2, W - 2.4, 0.6, f"— {s['author']}", 16, t["muted"], align="center")
        else:
            slide = self.new("FFFFFF")
            self.chrome(slide, s["title"])
            if layout == "two_columns":
                for i, (head, body) in enumerate(((s.get("left_title"), s["left"]), (s.get("right_title"), s["right"]))):
                    x = 0.6 + i * 6.2
                    self.box(slide, x, 1.6, 5.9, 5.0, t["light"])
                    if head:
                        self.text(slide, x + 0.3, 1.8, 5.3, 0.6, head, 20, t["accent"], True, t["heading"])
                    self.bullets(slide, x + 0.3, 2.5 if head else 1.9, 5.3, 4.0, body, 16)
            elif layout == "table" and s.get("table"):
                self.table(slide, s["table"], 0.6, 1.6, W - 1.2, 5.2)
            elif layout == "chart" and s.get("chart"):
                self.chart(slide, s["chart"], 0.6, 1.5, W - 1.2, 5.4)
            elif layout == "kpi" and s.get("items"):
                self.kpis(slide, s["items"])
            else:
                if s.get("chart"):
                    self.bullets(slide, 0.6, 1.7, 5.6, 5.0, s["bullets"], 18)
                    self.chart(slide, s["chart"], 6.4, 1.5, 6.3, 5.3)
                else:
                    self.bullets(slide, 0.6, 1.7, W - 1.2, 5.2, s["bullets"] or [s.get("subtitle") or ""], 20)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    def save(self, path: Path) -> Path:
        slides = self.spec["slides"]
        if slides[0]["layout"] != "title":
            slides = [{"layout": "title", "title": self.spec["title"], "subtitle": self.spec.get("subtitle", ""), "bullets": [],
                       "left": [], "right": [], "items": [], "notes": ""}] + slides
        for s in slides:
            self.slide(s)
        self.prs.core_properties.title = self.spec["title"]
        self.prs.core_properties.author = self.spec.get("author") or "Jarvis"
        self.prs.save(str(path))
        return path


def write(doc: dict, path: Path, links: dict | None = None) -> Path:
    return Deck(doc, links).save(path)
