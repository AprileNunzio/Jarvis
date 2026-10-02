import re
from datetime import datetime
from pathlib import Path

from features.documents import themes
from features.documents.spec import number

FORMATS = {"currency": '#,##0.00 "€"', "percent": "0.0%", "integer": "#,##0", "number": "#,##0.00", "date": "DD/MM/YYYY"}
DATE_RE = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})$|^(\d{4})-(\d{2})-(\d{2})$")
FIRST_ROW = 4
NO_SUM = re.compile(r"prezz|unitari|tariff|aliquot|sconto|%|percent|media|tasso|iva|margine", re.I)
DATA_SHEET = "Dati grafici"


def parse_date(value):
    if isinstance(value, datetime):
        return value
    m = DATE_RE.match(str(value or "").strip())
    if not m:
        return value
    try:
        if m.group(4):
            return datetime(int(m.group(4)), int(m.group(5)), int(m.group(6)))
        year = int(m.group(3))
        return datetime(year + 2000 if year < 100 else year, int(m.group(2)), int(m.group(1)))
    except ValueError:
        return value


def cell_value(value, kind: str):
    if isinstance(value, str) and value.startswith("="):
        return value
    if kind in ("currency", "number", "integer", "percent"):
        n = number(value)
        if n is None:
            return value
        if kind == "percent" and abs(n) > 1.5:
            n = n / 100
        return int(n) if kind == "integer" else n
    if kind == "date":
        return parse_date(value)
    return value


def make_chart(spec: dict, ws, first_col: int, first_row: int, theme: dict):
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart
    from openpyxl.chart.series import SeriesLabel
    cats, series, kind = spec["categories"], spec["series"], spec["chart"]
    from openpyxl.styles import Font
    ws.cell(row=first_row, column=first_col, value="Dati del grafico").font = Font(bold=True, size=10, color=theme["muted"])
    header = first_row + 1
    ws.cell(row=header, column=first_col, value="Categoria")
    for j, s in enumerate(series):
        ws.cell(row=header, column=first_col + 1 + j, value=s["name"])
    for i, c in enumerate(cats):
        ws.cell(row=header + 1 + i, column=first_col, value=c)
        for j, s in enumerate(series):
            ws.cell(row=header + 1 + i, column=first_col + 1 + j, value=s["values"][i])
    last = header + len(cats)
    labels = Reference(ws, min_col=first_col, min_row=header + 1, max_row=last)
    if kind in ("pie", "doughnut"):
        chart = DoughnutChart() if kind == "doughnut" else PieChart()
        chart.add_data(Reference(ws, min_col=first_col + 1, min_row=header, max_row=last), titles_from_data=True)
        chart.set_categories(labels)
    elif kind == "scatter":
        chart = ScatterChart()
        from openpyxl.chart import Series
        for j in range(len(series)):
            ys = Reference(ws, min_col=first_col + 1 + j, min_row=header + 1, max_row=last)
            s = Series(ys, labels, title=series[j]["name"])
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            chart.series.append(s)
    else:
        chart = {"line": LineChart, "area": AreaChart}.get(kind, BarChart)()
        if kind in ("bar", "column"):
            chart.type = "bar" if kind == "bar" else "col"
            chart.gapWidth = 60
        chart.add_data(Reference(ws, min_col=first_col + 1, max_col=first_col + len(series), min_row=header, max_row=last),
                       titles_from_data=True)
        chart.set_categories(labels)
    for j, s in enumerate(chart.series):
        color = theme["chart"][j % len(theme["chart"])]
        if kind not in ("pie", "doughnut"):
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color
        if not s.tx:
            s.tx = SeriesLabel(v=series[j]["name"])
    if kind in ("pie", "doughnut"):
        from openpyxl.chart.series import DataPoint
        for i in range(len(cats)):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = theme["chart"][i % len(theme["chart"])]
            chart.series[0].dPt.append(point)
    chart.title = spec.get("title") or None
    chart.style = 10
    chart.height, chart.width = 8.5, 17
    if len(series) == 1 and kind not in ("pie", "doughnut"):
        chart.legend = None
    return chart


def write_sheet(ws, sheet: dict, doc: dict, theme: dict, links: dict, data=None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    cols = list(sheet["columns"])
    formulas = dict(sheet.get("formulas") or {})
    for title in formulas:
        if title not in [c["title"] for c in cols]:
            cols.append({"title": title, "type": "number", "width": None})
    width = len(cols)
    thin = Side(style="thin", color="D5DBDB")
    head_fill = PatternFill("solid", fgColor=theme["primary"])
    zebra = PatternFill("solid", fgColor="F7F9FA")
    total_fill = PatternFill("solid", fgColor=theme["light"])
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=sheet["name"] if len(doc.get("sheets", [])) > 1 else doc["title"])
    ws.cell(row=1, column=1).font = Font(name=theme["heading"], size=18, bold=True, color=theme["primary"])
    ws.cell(row=2, column=1, value=sheet.get("description") or doc.get("subtitle") or "")
    ws.cell(row=2, column=1).font = Font(name=theme["body"], size=10, italic=True, color=theme["muted"])
    ws.row_dimensions[1].height = 30
    header = FIRST_ROW - 1
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=header, column=i, value=c["title"])
        cell.font = Font(name=theme["body"], bold=True, color=themes.readable_on(theme["primary"]), size=10.5)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=theme["accent"]))
    ws.row_dimensions[header].height = 24
    rows = sheet["rows"]
    for r_i, row in enumerate(rows):
        r = FIRST_ROW + r_i
        for c_i, c in enumerate(cols):
            if c["title"] in formulas:
                value = formulas[c["title"]].replace("{r}", str(r))
                value = value if value.startswith("=") else "=" + value
            else:
                value = cell_value(row[c_i] if c_i < len(row) else "", c["type"])
            cell = ws.cell(row=r, column=c_i + 1, value=value)
            cell.font = Font(name=theme["body"], size=10)
            cell.border = Border(bottom=thin)
            if c["type"] in FORMATS:
                cell.number_format = FORMATS[c["type"]]
            if r_i % 2:
                cell.fill = zebra
            cell.alignment = Alignment(vertical="center", horizontal="right" if c["type"] in FORMATS and c["type"] != "date" else "left")
    last = FIRST_ROW + len(rows) - 1
    if sheet.get("totals") and rows:
        tr = last + 1
        ws.cell(row=tr, column=1, value="Totale").font = Font(name=theme["body"], bold=True, size=10.5)
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=tr, column=i)
            cell.fill = total_fill
            cell.border = Border(top=Side(style="medium", color=theme["primary"]))
            if i > 1 and c["type"] in ("currency", "number", "integer") and not NO_SUM.search(c["title"]):
                letter = get_column_letter(i)
                cell.value = f"=SUM({letter}{FIRST_ROW}:{letter}{last})"
                cell.number_format = FORMATS[c["type"]]
                cell.font = Font(name=theme["body"], bold=True, size=10.5)
                cell.alignment = Alignment(horizontal="right")
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header}:{get_column_letter(width)}{last}"
    for i, c in enumerate(cols, start=1):
        longest = max([len(str(c["title"]))] + [len(str(row[i - 1])) for row in rows[:200] if i - 1 < len(row)])
        ws.column_dimensions[get_column_letter(i)].width = c.get("width") or max(10, min(48, longest + 4))
    below = last + (3 if sheet.get("totals") else 2)
    for li in sheet.get("links") or []:
        target = links.get(li.get("target", ""))
        if target:
            cell = ws.cell(row=below, column=1, value=f"→ {li['text'] or target}")
            cell.hyperlink = target
            cell.font = Font(name=theme["body"], color=theme["accent"], underline="single")
            below += 1
    if sheet.get("chart") and data is not None:
        start = data.max_row + 2 if data.max_row > 1 else 1
        chart = make_chart(sheet["chart"], data, 1, start, theme)
        data.cell(row=start, column=1).value = f"{ws.title} · {sheet['chart'].get('title') or 'grafico'}"
        ws.add_chart(chart, f"A{below + 2}")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header}:{header}"


def write(doc: dict, path: Path, links: dict | None = None) -> Path:
    from openpyxl import Workbook
    theme = themes.pick(doc.get("theme"), doc.get("palette"), doc.get("font"))
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    data = None
    if any(s.get("chart") for s in doc["sheets"]):
        data = wb.create_sheet(DATA_SHEET)
        data.sheet_properties.tabColor = theme["muted"]
        data.sheet_state = "hidden"
    for sheet in doc["sheets"]:
        name = re.sub(r"[\[\]:*?/\\]", "", sheet["name"])[:31] or "Foglio"
        base, n = name, 2
        while name.lower() in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name.lower())
        ws = wb.create_sheet(name, index=len(used) - 1)
        ws.sheet_properties.tabColor = theme["accent"]
        write_sheet(ws, sheet, doc, theme, links or {}, data)
    if data is not None:
        data.column_dimensions["A"].width = 28
        wb.move_sheet(data, offset=len(wb.sheetnames) - 1 - wb.sheetnames.index(DATA_SHEET))
    wb.active = 0
    wb.properties.title, wb.properties.creator = doc["title"], doc.get("author") or "Jarvis"
    wb.save(str(path))
    return path
